import json
from pathlib import Path
from docx import Document

PASTA = Path("perfis")
ARQUIVO_SAIDA = "perfis_analise_editais.docx"

doc = Document()

doc.add_heading("Perfis profissionais — Análise de Editais", level=0)

for arquivo in sorted(PASTA.glob("*.json")):
    with open(arquivo, "r", encoding="utf-8") as f:
        perfil = json.load(f)

    doc.add_heading(perfil.get("nome", arquivo.stem), level=1)

    campos = [
        ("Descrição", "descricao"),
        ("Graduações", "graduacoes"),
        ("Ferramentas", "ferramentas"),
        ("Áreas de interesse", "areas_interesse"),
        ("Idiomas", "idiomas"),
        ("Valor mínimo", "valor_minimo"),
        ("Experiência (anos)", "experiencia_anos"),
        ("Tem pós-graduação", "tem_pos_graduacao"),
        ("Tem mestrado", "tem_mestrado"),
        ("Tem doutorado", "tem_doutorado"),
        ("Contexto", "contexto"),
    ]

    for titulo, chave in campos:
        if chave not in perfil:
            continue

        valor = perfil[chave]

        doc.add_heading(titulo, level=2)

        if isinstance(valor, list):
            for item in valor:
                doc.add_paragraph(str(item), style="List Bullet")
        elif isinstance(valor, bool):
            doc.add_paragraph("Sim" if valor else "Não")
        elif chave == "valor_minimo":
            doc.add_paragraph(f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
        else:
            doc.add_paragraph(str(valor))

    doc.add_page_break()

doc.save(ARQUIVO_SAIDA)

print(f"Arquivo criado: {ARQUIVO_SAIDA}")
